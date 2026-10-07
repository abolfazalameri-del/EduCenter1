import sys, json, zipfile
import xml.dom.minidom as md
path = sys.argv[1]
out = {'ok': True, 'errors': []}
try:
    z = zipfile.ZipFile(path)
    bad = z.testzip()
    if bad: out['errors'].append('corrupt member ' + bad)
    for n in z.namelist():
        if n.endswith('.xml') or n.endswith('.rels'):
            try: md.parseString(z.read(n))
            except Exception as e: out['errors'].append('xml ' + n + ': ' + str(e))
    out['parts'] = z.namelist()
    if path.endswith('.docx'):
        import docx
        d = docx.Document(path)
        out['paragraphs'] = [p.text for p in d.paragraphs]
        out['tables'] = [[[c.text for c in r.cells] for r in t.rows] for t in d.tables]
        out['images'] = len(d.inline_shapes)
        out['bidi_paragraphs'] = sum(1 for p in d.paragraphs if p._p.pPr is not None and p._p.pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bidi') is not None)
        sec = d.sections[0]
        out['landscape'] = sec.page_width > sec.page_height
        out['footer_has_page_field'] = 'PAGE' in z.read('word/footer1.xml').decode('utf8')
        out['bold_runs'] = sum(1 for p in d.paragraphs for r in p.runs if r.bold)
    elif path.endswith('.xlsx'):
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        out['sheet'] = ws.title
        out['rtl'] = bool(ws.sheet_view.rightToLeft)
        out['freeze'] = ws.freeze_panes
        out['merged'] = [str(m) for m in ws.merged_cells.ranges]
        out['cells'] = [[(c.value if not isinstance(c.value, float) else c.value) for c in row] for row in ws.iter_rows()]
        out['types'] = [[(type(c.value).__name__) for c in row] for row in ws.iter_rows()]
        out['bold'] = [[bool(c.font.b) for c in row] for row in ws.iter_rows()]
        out['numfmt'] = [[c.number_format for c in row] for row in ws.iter_rows()]
        out['widths'] = {k: v.width for k, v in ws.column_dimensions.items()}
        out['orientation'] = ws.page_setup.orientation
        out['filter'] = ws.auto_filter.ref
except Exception as e:
    out['ok'] = False; out['errors'].append(repr(e))
out['ok'] = out['ok'] and not out['errors']
print(json.dumps(out, ensure_ascii=False))
